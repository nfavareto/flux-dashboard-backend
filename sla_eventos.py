"""
Pipeline de SLA de 1a respuesta e interacciones reales — requiere el
dump de S1 a NIVEL EVENTO (columnas: case_id, channel, user_id,
queued_at, sa, ht, group_name, Usuario), distinto del case_details
(nivel caso). Se cruza por case_id con case_details para saber que
casos son del Contact Center.
"""
from collections import Counter, defaultdict
from datetime import datetime

import openpyxl

from .common import norm_user_id, CC_USERS
from .s1 import load_case_details, _bucket


def build_case_bucket_map(case_details_path: str) -> dict:
    rows, idx, g = load_case_details(case_details_path)
    return {g(r, 'case_id'): _bucket(g(r, 'user_id_ev')) for r in rows}


def process_sla_mes(path_evento: str, case_bucket: dict, mes_key: str):
    """Devuelve (sla_1a_respuesta_pct_cc, interacciones_totales_cc, interacciones_por_agente_cc)
    para los eventos cuyo queued_at cae en mes_key."""
    wb = openpyxl.load_workbook(path_evento, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(min_row=1, values_only=True)
    hdr = list(next(it))
    idx = {h: i for i, h in enumerate(hdr)}

    def g(r, n):
        i = idx.get(n)
        return r[i] if i is not None and i < len(r) else None

    first = {}  # case_id -> min sa (primera respuesta no-bot, sa>0)
    inter_by_agent = Counter()
    n_inter_cc = 0

    for r in ws.iter_rows(min_row=2, values_only=True):
        cid = g(r, 'case_id')
        if cid is None:
            continue
        q = g(r, 'queued_at')
        if not (isinstance(q, datetime) and q.strftime('%Y-%m') == mes_key):
            continue
        uid = norm_user_id(g(r, 'user_id'))
        gname = str(g(r, 'group_name') or '')
        is_bot = 'bot' in gname.lower()

        if uid in CC_USERS:
            inter_by_agent[uid] += 1
            n_inter_cc += 1

        if not is_bot:
            sa = g(r, 'sa')
            try:
                saf = float(sa)
            except Exception:
                saf = None
            if saf is not None and saf > 0:
                if cid not in first or saf < first[cid]:
                    first[cid] = saf

    vals_cc = [sa for cid, sa in first.items() if case_bucket.get(cid) == 'cc']
    sla_pct = round(sum(1 for x in vals_cc if x <= 180) / len(vals_cc) * 100, 1) if vals_cc else None

    return sla_pct, n_inter_cc, dict(inter_by_agent)
