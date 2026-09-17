"""
Pipeline de S1 Digital — nivel caso (case_details).

Regla de negocio (definida por el cliente): un caso es del CONTACT CENTER
si lo resolvio uno de los 9 agentes del equipo (CC_USERS) o el bot
(user_id_ev vacio, porque el bot solo deriva a esos agentes). Cualquier
otro usuario resolutor => Sucursal.

Genera, para 'cc' y 'suc' por separado:
  - S1C / S1SUC   (casos, resolucion, tiempos, buckets, grupos, campanas)
  - S1DAY / S1DAYSUC (recibidos/en cola/resueltos por dia)
  - WAD.cc / WAD.suc  (whatsapp entrante/saliente)
  - RESHA (solo CC: resueltos por humano vs bot)
"""
from collections import Counter, defaultdict
from datetime import datetime

import openpyxl

from .common import fix_mojibake, norm_user_id, CC_USERS


def _bucket(user_id_ev):
    uev = norm_user_id(user_id_ev)
    if uev is None:
        return 'cc'  # bot: solo deriva a Contact Center
    return 'cc' if uev in CC_USERS else 'suc'


def _channel(c):
    c = (c or '').lower()
    if 'whats' in c:
        return 'wa'
    if 'mail' in c:
        return 'email'
    if 'face' in c:
        return 'fb'
    if 'insta' in c:
        return 'ig'
    return 'otro'


def load_case_details(path: str):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(min_row=1, values_only=True)
    hdr = list(next(it))
    idx = {h: i for i, h in enumerate(hdr)}

    def g(r, n):
        i = idx.get(n)
        return r[i] if i is not None and i < len(r) else None

    rows = [r for r in it if g(r, 'case_id') is not None]
    wb.close()
    return rows, idx, g


def _build_group(rows, g):
    n = len(rows)
    res = sum(1 for r in rows if g(r, 'case_state') == 5)
    sal = sum(1 for r in rows if g(r, 'case_out') == 1)
    ch = Counter(_channel(g(r, 'channel')) for r in rows)

    rts = []
    for r in rows:
        d, rr = g(r, 'date'), g(r, 'resolution_at')
        if isinstance(d, datetime) and isinstance(rr, datetime):
            m = (rr - d).total_seconds() / 60
            if m >= 0:
                rts.append(m)
    rts.sort()
    N = len(rts)
    med = rts[N // 2] if N else 0
    avg = sum(rts) / N if N else 0

    def pc(T):
        return round(sum(1 for x in rts if x <= T) / N * 100, 1) if N else 0

    grp_c, grp_res, grp_rt = Counter(), Counter(), defaultdict(list)
    for r in rows:
        gn = g(r, 'group_name') or 'Sin grupo'
        grp_c[gn] += 1
        if g(r, 'case_state') == 5:
            grp_res[gn] += 1
        d, rr = g(r, 'date'), g(r, 'resolution_at')
        if isinstance(d, datetime) and isinstance(rr, datetime):
            m = (rr - d).total_seconds() / 60
            if m >= 0:
                grp_rt[gn].append(m)
    grp = []
    for gn, c in grp_c.most_common(10):
        rr = sorted(grp_rt[gn])
        gm = rr[len(rr) // 2] if rr else 0
        grp.append({'g': fix_mojibake(gn), 'casos': c,
                     'res': round(grp_res[gn] / c * 100, 1) if c else 0, 'med': round(gm)})

    camp = Counter(fix_mojibake(g(r, 'campaign_name') or 'Sin campana') for r in rows)

    return {
        'casos': n, 'res': res, 'res_pct': round(res / n * 100, 1) if n else 0,
        'sal': sal, 'ent': n - sal, 'med': round(med), 'avg': round(avg),
        'buckets': {'h1': pc(60), 'h2': pc(120), 'h4': pc(240), 'h24': pc(1440)},
        'wa': ch.get('wa', 0), 'email': ch.get('email', 0), 'fb': ch.get('fb', 0), 'ig': ch.get('ig', 0),
        'grp': grp, 'camp': [{'c': c, 'n': n2} for c, n2 in camp.most_common(8)],
    }


def _daily(rows, g):
    dd = defaultdict(lambda: {'recib': 0, 'cola': 0, 'resuel': 0})
    for r in rows:
        d = g(r, 'date')
        if not isinstance(d, datetime):
            continue
        k = d.strftime('%Y-%m-%d')
        x = dd[k]
        x['recib'] += 1
        if 'bot' not in str(g(r, 'group_name') or '').lower():
            x['cola'] += 1
        if g(r, 'case_state') == 5:
            x['resuel'] += 1
    return [{'d': int(k[-2:]), 'recib': v['recib'], 'cola': v['cola'], 'resuel': v['resuel']}
            for k, v in sorted(dd.items())]


def _wad(rows, g):
    a = defaultdict(int)
    for r in rows:
        c = _channel(g(r, 'channel'))
        out = (g(r, 'case_out') == 1)
        a[c + ('_out' if out else '_in')] += 1
        a[c] += 1
    return {'wa_in': a['wa_in'], 'wa_out': a['wa_out'], 'wa': a['wa'],
            'email': a['email'], 'fb': a['fb'], 'ig': a['ig'], 'email_out': a['email_out']}


def process_case_details(path: str, mes_key: str) -> dict:
    """Devuelve todo lo que hace falta actualizar para un mes de S1:
    {s1c_cc, s1c_suc, day_cc, day_suc, wad_cc, wad_suc, resha}"""
    rows, idx, g = load_case_details(path)
    cc_rows = [r for r in rows if _bucket(g(r, 'user_id_ev')) == 'cc']
    suc_rows = [r for r in rows if _bucket(g(r, 'user_id_ev')) == 'suc']

    hum = sum(1 for r in cc_rows if norm_user_id(g(r, 'user_id_ev')) in CC_USERS and g(r, 'case_state') == 5)
    auto = sum(1 for r in cc_rows if norm_user_id(g(r, 'user_id_ev')) is None and g(r, 'case_state') == 5)

    return {
        's1c_cc': _build_group(cc_rows, g),
        's1c_suc': _build_group(suc_rows, g),
        'day_cc': _daily(cc_rows, g),
        'day_suc': _daily(suc_rows, g),
        'wad_cc': _wad(cc_rows, g),
        'wad_suc': _wad(suc_rows, g),
        'resha': {'res': hum + auto, 'humano': hum, 'auto': auto},
        'mes': mes_key,
    }
