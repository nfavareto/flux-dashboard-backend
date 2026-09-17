"""
Pipeline de Presentismo — ingreso/egreso por agente y dia, cruzado
con el turno ASIGNADO (grilla de turnos) para detectar llegadas tarde
reales y cambios de turno.
"""
import re
import unicodedata
from datetime import time, datetime

import openpyxl


def _toks(s):
    s = unicodedata.normalize('NFKD', str(s or '')).encode('ascii', 'ignore').decode('ascii').lower()
    return set(t for t in s.replace('.', ' ').split() if len(t) > 2)


def load_turnomap(path_grilla: str) -> dict:
    """Presentismo_Flux.xlsx: una hoja por mes (ENERO 26, FEBRERO 26, ...),
    con bloques TM (manana) / TT (tarde) y una fila por agente debajo.
    Devuelve {mes_key: [(tokens_nombre, 'manana'|'tarde'), ...]}."""
    MESN = {'ENERO': 1, 'FEBRERO': 2, 'MARZO': 3, 'ABRIL': 4, 'MAYO': 5, 'JUNIO': 6,
            'JULIO': 7, 'AGOSTO': 8, 'SEPTIEMBRE': 9, 'OCTUBRE': 10, 'NOVIEMBRE': 11, 'DICIEMBRE': 12}
    STOP = ('total', 'horas a cubrir', 'diferencia', 'horas extras')
    wb = openpyxl.load_workbook(path_grilla, data_only=True)
    out = {}
    for sh in wb.sheetnames:
        m = re.match(r'([A-Z]+)\s*(\d{2})', sh.strip())
        if not m or m.group(1) not in MESN:
            continue
        mes_key = f"20{m.group(2)}-{MESN[m.group(1)]:02d}"
        ws = wb[sh]
        cur = None
        entries = []
        for r in ws.iter_rows(values_only=True):
            a = r[0]
            if not a:
                continue
            au = str(a).strip()
            al = au.lower()
            if any(al.startswith(s) for s in STOP):
                break
            if au.upper().startswith('TM'):
                cur = 'manana'
                continue
            if au.upper().startswith('TT'):
                cur = 'tarde'
                continue
            if cur and len(au.split()) >= 2 and any(ch.isalpha() for ch in au):
                entries.append((sorted(_toks(au)), cur))
        out[mes_key] = entries
    return out


def _match_turno(nombre, mes_key, turnomap):
    nt = _toks(nombre)
    best, bn = None, 0
    for tk, tu in turnomap.get(mes_key, []):
        ov = len(nt & set(tk))
        if ov > bn:
            bn, best = ov, tu
    return best


def _hhmm(v):
    if isinstance(v, (time, datetime)):
        return f"{v.hour:02d}:{v.minute:02d}"
    m = re.search(r'(\d{1,2}):(\d{2})', str(v or ''))
    return f"{int(m.group(1)):02d}:{m.group(2)}" if m else None


def _tomin(hm):
    if not hm:
        return None
    h, m = hm.split(':')
    return int(h) * 60 + int(m)


def process_presentismo(path_reporte_usuarios: str, mes_key: str, u2n: dict, u2r: dict, turnomap: dict):
    """path_reporte_usuarios: mismo formato que Presentismo_Semestre.xlsx / ReporteUsuarios:
    columnas Usuario, Año, Mes, Dia, Hora Primer Registro, Hora Ultimo Registro, Tiempo Total.
    u2n: {user_id: nombre_completo} de los agentes del Contact Center (el archivo trae TODOS
    los usuarios de la plataforma, incluidas sucursales; solo se procesan los que estan en u2n).
    u2r: {user_id: rol}. Devuelve (lista_de_registros_del_mes, pre_stats_del_mes)."""
    wb = openpyxl.load_workbook(path_reporte_usuarios, data_only=True, read_only=True)
    ws = wb['hoja 1'] if 'hoja 1' in wb.sheetnames else wb[wb.sheetnames[0]]
    it = ws.iter_rows(min_row=1, values_only=True)
    hdr = list(next(it))
    idx = {h: i for i, h in enumerate(hdr)}
    ykey = 'Año' if 'Año' in idx else next(h for h in idx if h and 'o' in str(h) and len(str(h)) <= 4)

    out = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        u = r[idx['Usuario']]
        if not u or u not in u2n:
            continue  # solo agentes del Contact Center (u2n = los 9 conocidos)
        y, mo, da = r[idx[ykey]], r[idx['Mes']], r[idx['Dia']]
        if not (y and mo and da):
            continue
        mk = f"{int(y):04d}-{int(mo):02d}"
        if mk != mes_key:
            continue
        fecha = f"{mk}-{int(da):02d}"
        ent = _hhmm(r[idx['Hora Primer Registro']])
        sal = _hhmm(r[idx['Hora Ultimo Registro']])
        tot = _hhmm(r[idx['Tiempo Total']])
        nombre = u2n.get(u, str(u).capitalize())
        ta = _match_turno(nombre, mk, turnomap)
        em = _tomin(ent)
        treal = 'manana' if (em is not None and em < 13 * 60) else 'tarde'
        turno = ta or treal
        esp = '09:00' if turno == 'manana' else '13:30'
        espm = _tomin(esp)

        estado, tarde, tardem = 'ok', False, 0
        if ta and treal != ta:
            estado = 'cambio'
        elif em is not None and espm is not None and (em - espm) > 5:
            estado, tarde, tardem = 'tarde', True, em - espm

        out.append({
            'user': u, 'nombre': nombre, 'fecha': fecha, 'hora': ent or '--',
            'hora_entrada': ent or '--', 'hora_salida': sal or '--', 'tiempo_total': tot or '--',
            'turno': turno, 'hora_esperada': esp, 'estado': estado, 'tarde': tarde,
            'minutos_tarde': tardem, 'rol': u2r.get(u, 'operador'),
        })

    stats = {
        'total': len(out),
        'ok': sum(1 for x in out if x['estado'] == 'ok'),
        'tarde': sum(1 for x in out if x['estado'] == 'tarde'),
        'cambio': sum(1 for x in out if x['estado'] == 'cambio'),
    }
    return out, stats
