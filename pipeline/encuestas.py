"""
Pipeline de Encuestas — IVR de encuesta post-llamada (3 preguntas, 1-5).
Se cruza con la tipificacion de la llamada via Reporte Gestiones
(survey.id_gestion <-> gestiones.ID Gestion).
"""
import re
import zipfile
from collections import Counter, defaultdict
from xml.etree import ElementTree as ET

import openpyxl

from .common import fix_mojibake

_NS = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'


def _score(v):
    try:
        s = int(float(v))
        return s if 1 <= s <= 5 else None
    except Exception:
        return None


def _read_xlsx_raw(path):
    """Lector manual via XML: algunos exports de Gestiones traen un
    stylesheet corrupto que openpyxl no puede abrir (Fill() sin argumentos).
    Este fallback ignora estilos y lee solo los valores."""
    z = zipfile.ZipFile(path)
    ss = []
    if 'xl/sharedStrings.xml' in z.namelist():
        r = ET.fromstring(z.read('xl/sharedStrings.xml'))
        for si in r.findall(_NS + 'si'):
            ss.append(''.join(t.text or '' for t in si.iter(_NS + 't')))
    sheet = sorted(n for n in z.namelist() if re.match(r'xl/worksheets/sheet\d+\.xml', n))[0]
    root = ET.fromstring(z.read(sheet))
    grid = []
    for row in root.iter(_NS + 'row'):
        cells = []
        for c in row.findall(_NS + 'c'):
            t = c.get('t')
            v = c.find(_NS + 'v')
            val = ''
            if v is not None:
                val = v.text
                if t == 's':
                    val = ss[int(val)]
            cells.append(val)
        grid.append(cells)
    hdr = grid[0]
    return hdr, [{hdr[i]: (r[i] if i < len(r) else '') for i in range(len(hdr))} for r in grid[1:]]


def load_gestiones_idg2cat(path_gestiones: str) -> dict:
    try:
        wb = openpyxl.load_workbook(path_gestiones, data_only=True, read_only=True)
        ws = wb[wb.sheetnames[0]]
        it = ws.iter_rows(min_row=1, values_only=True)
        hdr = list(next(it))
        idx = {h: i for i, h in enumerate(hdr)}
        out = {}
        for r in it:
            k = r[idx['ID Gestion']] if 'ID Gestion' in idx else None
            c = r[idx['Categoria']] if 'Categoria' in idx else None
            if k and c:
                out[str(k)] = str(c).strip()
        wb.close()
        return out
    except TypeError:
        # Fallback: stylesheet corrupto -> leer via XML crudo
        _, rows = _read_xlsx_raw(path_gestiones)
        out = {}
        for x in rows:
            k = x.get('ID Gestion')
            c = (x.get('Categoria') or '').strip()
            if k and c:
                out[str(k)] = c
        return out


def process_encuesta_mes(path_ivr_encuesta: str, idg2cat: dict, mes_key: str, td_atendidas: int = None):
    """Devuelve (bymonth_entry, bytip_list) para el mes_key."""
    wb = openpyxl.load_workbook(path_ivr_encuesta, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(min_row=1, values_only=True)
    hdr = list(next(it))
    idx = {h: i for i, h in enumerate(hdr)}

    def g(r, n):
        i = idx.get(n)
        return r[i] if i is not None and i < len(r) else None

    rows = [r for r in it if g(r, 'call_id') is not None and str(g(r, 'Fecha Inicio'))[:7] == mes_key]

    n = len(rows)
    resp = 0
    p1, p2, p3 = [], [], []
    by_cat = defaultdict(lambda: {'n': 0, 'resp': 0, 'p1': [], 'p2': [], 'p3': []})

    for r in rows:
        s1, s2, s3 = _score(g(r, 'Step_2')), _score(g(r, 'Step_3')), _score(g(r, 'Step_4'))
        has_any = bool(s1 or s2 or s3)
        if has_any:
            resp += 1
        for s, lst in ((s1, p1), (s2, p2), (s3, p3)):
            if s:
                lst.append(s)
        cat = idg2cat.get(str(g(r, 'id_gestion')))
        if cat:
            a = by_cat[cat]
            a['n'] += 1
            if has_any:
                a['resp'] += 1
            for s, k in ((s1, 'p1'), (s2, 'p2'), (s3, 'p3')):
                if s:
                    a[k].append(s)

    def metrics(lst_p1, lst_p2, lst_p3):
        def top(lst):
            return round(sum(1 for x in lst if x >= 4) / len(lst) * 100, 1) if lst else 0

        def avg(lst):
            return round(sum(lst) / len(lst), 2) if lst else 0

        def nps(lst):
            if not lst:
                return 0
            prom = sum(1 for x in lst if x == 5)
            det = sum(1 for x in lst if x <= 3)
            return round((prom - det) / len(lst) * 100)

        return {
            'fcr': top(lst_p1), 'fcr_avg': avg(lst_p1), 'n1': len(lst_p1),
            'csat': top(lst_p2), 'csat_avg': avg(lst_p2), 'n2': len(lst_p2),
            'nps': nps(lst_p3), 'nps_avg': avg(lst_p3), 'n3': len(lst_p3),
        }

    m = metrics(p1, p2, p3)
    bymonth = {
        'n': n, 'resp': resp, 'rate': round(resp / n * 100, 1) if n else 0,
        'deriv': n, 'atend': td_atendidas,
        **m,
    }

    bytip = []
    for cat, a in by_cat.items():
        if a['n'] < 5:
            continue
        mm = metrics(a['p1'], a['p2'], a['p3'])
        bytip.append({'cat': fix_mojibake(cat), 'n': a['n'], **mm})
    bytip.sort(key=lambda x: -x['n'])

    return bymonth, bytip
